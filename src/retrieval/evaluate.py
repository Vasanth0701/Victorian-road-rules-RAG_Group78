from pathlib import Path
import csv
import json
import re

import faiss
import numpy as np
import openpyxl
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer


# ============================================================
# CONFIGURATION
# ============================================================

CHUNKS_FILE = Path("data/processed/chunks.jsonl")

GROUND_TRUTH_FILE = Path(
    "data/test_questions/Road_Rules_Ground_Truth_Final_10_Questions.xlsx"
)

FAISS_INDEX_FILE = Path("indexes/road_rules.faiss")

RESULTS_DIR = Path("results")
CSV_OUTPUT_FILE = RESULTS_DIR / "retrieval_evaluation.csv"

MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

TOP_K_VALUES = [1, 3, 5]


# ============================================================
# LOAD CHUNKS
# ============================================================

def load_chunks():
    chunks = []

    with CHUNKS_FILE.open("r", encoding="utf-8") as file:
        for line in file:
            line = line.strip()

            if line:
                chunks.append(json.loads(line))

    return chunks


# ============================================================
# LOAD GROUND TRUTH
# ============================================================

def load_ground_truth():
    workbook = openpyxl.load_workbook(
        GROUND_TRUTH_FILE,
        data_only=True
    )

    sheet = workbook["Ground Truth"]

    headers = [
        cell.value
        for cell in next(
            sheet.iter_rows(min_row=1, max_row=1)
        )
    ]

    question_id_col = headers.index("Question ID")
    question_col = headers.index("User Question")
    source_col = headers.index("Expected Source ID")

    questions = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        question_id = row[question_id_col]
        question = row[question_col]
        expected_source = row[source_col]

        if question_id and question and expected_source:
            questions.append(
                {
                    "question_id": str(question_id).strip(),
                    "question": str(question).strip(),
                    "expected_source_id": str(expected_source).strip(),
                }
            )

    return questions


# ============================================================
# TOKENIZATION
# ============================================================

def tokenize(text):
    return re.findall(
        r"\b\w+\b",
        text.lower()
    )


# ============================================================
# BM25
# ============================================================

def build_bm25(chunks):
    documents = [
        f"{chunk['section']} {chunk['text']}"
        for chunk in chunks
    ]

    tokenized_documents = [
        tokenize(document)
        for document in documents
    ]

    return BM25Okapi(tokenized_documents)


def bm25_retrieve(
    query,
    bm25,
    chunks,
    top_k
):
    scores = bm25.get_scores(
        tokenize(query)
    )

    ranked_indices = np.argsort(
        scores
    )[::-1][:top_k]

    results = []

    for index_id in ranked_indices:
        results.append(chunks[index_id])

    return results


# ============================================================
# DENSE FAISS
# ============================================================

def dense_retrieve(
    query,
    model,
    index,
    chunks,
    top_k
):
    query_embedding = model.encode(
        [query],
        normalize_embeddings=True
    )

    query_embedding = np.asarray(
        query_embedding,
        dtype="float32"
    )

    _, indices = index.search(
        query_embedding,
        top_k
    )

    results = []

    for index_id in indices[0]:
        if index_id >= 0:
            results.append(
                chunks[index_id]
            )

    return results


# ============================================================
# EVALUATION HELPERS
# ============================================================

def is_hit(
    results,
    expected_source
):
    return any(
        result["source_id"] == expected_source
        for result in results
    )


def get_source_ids(results):
    return [
        result["source_id"]
        for result in results
    ]


def find_first_correct_rank(
    results,
    expected_source
):
    for rank, result in enumerate(
        results,
        start=1
    ):
        if result["source_id"] == expected_source:
            return rank

    return None


# ============================================================
# MAIN EVALUATION
# ============================================================

def evaluate():

    # --------------------------------------------------------
    # Load data and retrieval systems
    # --------------------------------------------------------

    print("Loading chunks...")
    chunks = load_chunks()

    print("Loading ground truth...")
    questions = load_ground_truth()

    print("Loading BM25...")
    bm25 = build_bm25(chunks)

    print("Loading embedding model...")
    model = SentenceTransformer(
        MODEL_NAME
    )

    print("Loading FAISS index...")
    index = faiss.read_index(
        str(FAISS_INDEX_FILE)
    )

    print()
    print(f"Chunks: {len(chunks)}")
    print(
        f"Ground-truth questions: "
        f"{len(questions)}"
    )
    print()


    # --------------------------------------------------------
    # Metric storage
    # --------------------------------------------------------

    methods = [
        "BM25",
        "Dense FAISS"
    ]

    totals = {
        method: {
            k: 0
            for k in TOP_K_VALUES
        }
        for method in methods
    }

    evaluation_rows = []


    # --------------------------------------------------------
    # Evaluate every question
    # --------------------------------------------------------

    for item in questions:

        question_id = item["question_id"]
        question = item["question"]
        expected_source = (
            item["expected_source_id"]
        )

        print("=" * 70)

        print(
            f"{question_id}: "
            f"{question}"
        )

        print(
            f"Expected source: "
            f"{expected_source}"
        )

        print()


        # ----------------------------------------------------
        # Retrieve Top-5 once
        # ----------------------------------------------------

        bm25_top5 = bm25_retrieve(
            question,
            bm25,
            chunks,
            5
        )

        dense_top5 = dense_retrieve(
            question,
            model,
            index,
            chunks,
            5
        )


        # ----------------------------------------------------
        # Evaluate Top-1, Top-3 and Top-5
        # ----------------------------------------------------

        for k in TOP_K_VALUES:

            bm25_results = bm25_top5[:k]
            dense_results = dense_top5[:k]

            bm25_hit = is_hit(
                bm25_results,
                expected_source
            )

            dense_hit = is_hit(
                dense_results,
                expected_source
            )

            if bm25_hit:
                totals["BM25"][k] += 1

            if dense_hit:
                totals["Dense FAISS"][k] += 1

            bm25_sources = get_source_ids(
                bm25_results
            )

            dense_sources = get_source_ids(
                dense_results
            )

            print(
                f"Top-{k:<2} | "
                f"BM25: {bm25_sources} "
                f"{'HIT' if bm25_hit else 'MISS'}"
            )

            print(
                f"       | "
                f"Dense: {dense_sources} "
                f"{'HIT' if dense_hit else 'MISS'}"
            )


        # ----------------------------------------------------
        # First correct source rank
        # ----------------------------------------------------

        bm25_correct_rank = (
            find_first_correct_rank(
                bm25_top5,
                expected_source
            )
        )

        dense_correct_rank = (
            find_first_correct_rank(
                dense_top5,
                expected_source
            )
        )


        # ----------------------------------------------------
        # Save detailed result for CSV
        # ----------------------------------------------------

        evaluation_rows.append(
            {
                "question_id":
                    question_id,

                "question":
                    question,

                "expected_source":
                    expected_source,

                "bm25_rank1":
                    bm25_top5[0]["source_id"],

                "bm25_top3":
                    "|".join(
                        get_source_ids(
                            bm25_top5[:3]
                        )
                    ),

                "bm25_top5":
                    "|".join(
                        get_source_ids(
                            bm25_top5
                        )
                    ),

                "bm25_correct_rank":
                    bm25_correct_rank,

                "bm25_hit1":
                    is_hit(
                        bm25_top5[:1],
                        expected_source
                    ),

                "bm25_hit3":
                    is_hit(
                        bm25_top5[:3],
                        expected_source
                    ),

                "bm25_hit5":
                    is_hit(
                        bm25_top5,
                        expected_source
                    ),

                "dense_rank1":
                    dense_top5[0]["source_id"],

                "dense_top3":
                    "|".join(
                        get_source_ids(
                            dense_top5[:3]
                        )
                    ),

                "dense_top5":
                    "|".join(
                        get_source_ids(
                            dense_top5
                        )
                    ),

                "dense_correct_rank":
                    dense_correct_rank,

                "dense_hit1":
                    is_hit(
                        dense_top5[:1],
                        expected_source
                    ),

                "dense_hit3":
                    is_hit(
                        dense_top5[:3],
                        expected_source
                    ),

                "dense_hit5":
                    is_hit(
                        dense_top5,
                        expected_source
                    ),
            }
        )

        print()


    # ========================================================
    # SAVE CSV
    # ========================================================

    RESULTS_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    fieldnames = [
        "question_id",
        "question",
        "expected_source",

        "bm25_rank1",
        "bm25_top3",
        "bm25_top5",
        "bm25_correct_rank",
        "bm25_hit1",
        "bm25_hit3",
        "bm25_hit5",

        "dense_rank1",
        "dense_top3",
        "dense_top5",
        "dense_correct_rank",
        "dense_hit1",
        "dense_hit3",
        "dense_hit5",
    ]

    with CSV_OUTPUT_FILE.open(
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames
        )

        writer.writeheader()

        writer.writerows(
            evaluation_rows
        )


    # ========================================================
    # FINAL HIT@K SUMMARY
    # ========================================================

    total_questions = len(questions)

    print("=" * 70)
    print("FINAL RETRIEVAL SUMMARY")
    print("=" * 70)

    for method in methods:

        print()
        print(method)

        for k in TOP_K_VALUES:

            hits = totals[method][k]

            percentage = (
                hits
                / total_questions
                * 100
                if total_questions
                else 0
            )

            print(
                f"Hit@{k}: "
                f"{hits}/{total_questions} "
                f"({percentage:.1f}%)"
            )


    # ========================================================
    # MRR@5
    # ========================================================

    bm25_reciprocal_ranks = []
    dense_reciprocal_ranks = []

    for row in evaluation_rows:

        bm25_rank = row[
            "bm25_correct_rank"
        ]

        dense_rank = row[
            "dense_correct_rank"
        ]

        if bm25_rank:
            bm25_reciprocal_ranks.append(
                1 / bm25_rank
            )
        else:
            bm25_reciprocal_ranks.append(0)

        if dense_rank:
            dense_reciprocal_ranks.append(
                1 / dense_rank
            )
        else:
            dense_reciprocal_ranks.append(0)


    bm25_mrr = np.mean(
        bm25_reciprocal_ranks
    )

    dense_mrr = np.mean(
        dense_reciprocal_ranks
    )


    print()
    print("=" * 70)
    print("MRR@5")
    print("=" * 70)

    print(
        f"BM25 MRR@5: "
        f"{bm25_mrr:.4f}"
    )

    print(
        f"Dense FAISS MRR@5: "
        f"{dense_mrr:.4f}"
    )

    print()

    print(
        f"CSV results saved to: "
        f"{CSV_OUTPUT_FILE}"
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    evaluate()
